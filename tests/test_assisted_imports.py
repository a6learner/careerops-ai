from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document

from careerops.adapters.document_text_extractor import (
    DocumentExtractionError,
    extract_document_text,
)
from careerops.adapters.fake_job_search import FakeJobSearchProvider
from careerops.domain.assisted_imports import (
    CvImportErrorCode,
    DocumentType,
    JobSearchQuery,
)
from careerops.domain.errors import InvalidDocumentError
from careerops.services.assisted_imports import (
    detect_document_type,
    extract_experience_drafts,
    validate_extracted_drafts,
    validate_search_results,
)


@pytest.mark.parametrize(
    ("filename", "content", "expected"),
    [
        ("resume.pdf", b"%PDF-1.7\n", DocumentType.PDF),
        ("resume.doc", bytes.fromhex("D0CF11E0A1B11AE1") + b"data", DocumentType.DOC),
        ("resume.rtf", b"{\\rtf1 CareerOps}", DocumentType.RTF),
        ("resume.txt", b"CareerOps AI", DocumentType.TEXT),
        ("resume.md", b"# CareerOps AI", DocumentType.MARKDOWN),
    ],
)
def test_detect_document_type_uses_content_and_extension(
    filename: str, content: bytes, expected: DocumentType
) -> None:
    assert detect_document_type(filename, content) == expected


def test_detect_document_type_rejects_mismatched_or_binary_text() -> None:
    with pytest.raises(InvalidDocumentError):
        detect_document_type("resume.pdf", b"not a pdf")
    with pytest.raises(InvalidDocumentError):
        detect_document_type("resume.txt", b"a\x00b")


def test_docx_extraction_and_experience_drafts_are_grounded() -> None:
    document = Document()
    document.add_heading("CareerOps AI", level=1)
    document.add_paragraph("Built FastAPI and PostgreSQL services.")
    stream = BytesIO()
    document.save(stream)
    content = stream.getvalue()

    assert detect_document_type("resume.docx", content) == DocumentType.DOCX
    text = extract_document_text(DocumentType.DOCX, "resume.docx", content)
    drafts = extract_experience_drafts(text)
    validate_extracted_drafts(text, drafts)

    assert drafts[0].title == "CareerOps AI"
    assert "FastAPI" in drafts[0].description
    assert drafts[0].source_refs[0].quote in text


def test_archive_expansion_limit_rejects_suspicious_docx() -> None:
    stream = BytesIO()
    with ZipFile(stream, "w", ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "x" * (21 * 1024 * 1024))

    with pytest.raises(DocumentExtractionError) as raised:
        extract_document_text(DocumentType.DOCX, "resume.docx", stream.getvalue())

    assert raised.value.code == CvImportErrorCode.INVALID_ARCHIVE


def test_malformed_pdf_becomes_a_durable_extraction_failure() -> None:
    with pytest.raises(DocumentExtractionError) as raised:
        extract_document_text(DocumentType.PDF, "resume.pdf", b"%PDF-1.7\nnot a pdf")

    assert raised.value.code == CvImportErrorCode.EXTRACTOR_FAILED


@pytest.mark.asyncio
async def test_fake_search_is_deterministic_and_marks_incomplete_results() -> None:
    provider = FakeJobSearchProvider()
    query = JobSearchQuery(
        query="Backend Engineer",
        location="Los Angeles",
        remote=False,
        published_within_days=14,
    )

    first = await provider.search(query)
    second = await provider.search(query)
    validate_search_results(first)

    assert first == second
    assert first[0].raw_description is not None
    assert first[1].raw_description is None
