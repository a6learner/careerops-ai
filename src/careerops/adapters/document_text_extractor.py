import os
import shutil
import subprocess
import tempfile
import zipfile
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from careerops.domain.assisted_imports import CvImportErrorCode, DocumentType

MAX_ARCHIVE_MEMBERS = 1_000
MAX_ARCHIVE_EXPANDED_BYTES = 20 * 1024 * 1024
MAX_ARCHIVE_RATIO = 100
MAX_NORMALIZED_TEXT_LENGTH = 200_000


class DocumentExtractionError(Exception):
    def __init__(self, code: CvImportErrorCode) -> None:
        super().__init__(code.value)
        self.code = code


def extract_document_text(document_type: DocumentType, filename: str, content: bytes) -> str:
    try:
        if document_type == DocumentType.PDF:
            text = _extract_pdf(content)
        elif document_type == DocumentType.DOCX:
            _validate_archive(content)
            text = _extract_docx(content)
        elif document_type in {DocumentType.DOC, DocumentType.ODT, DocumentType.RTF}:
            if document_type == DocumentType.ODT:
                _validate_archive(content)
            text = _extract_with_libreoffice(filename, content)
        else:
            text = content.decode("utf-8-sig")
    except DocumentExtractionError:
        raise
    except Exception as exc:
        raise DocumentExtractionError(CvImportErrorCode.EXTRACTOR_FAILED) from exc

    normalized = "\n".join(line.rstrip() for line in text.replace("\r\n", "\n").split("\n"))
    normalized = normalized.strip()
    if not normalized:
        raise DocumentExtractionError(CvImportErrorCode.EMPTY_TEXT)
    if len(normalized) > MAX_NORMALIZED_TEXT_LENGTH:
        raise DocumentExtractionError(CvImportErrorCode.EXTRACTOR_FAILED)
    return normalized


def _extract_pdf(content: bytes) -> str:
    reader = PdfReader(BytesIO(content))
    if reader.is_encrypted:
        raise DocumentExtractionError(CvImportErrorCode.ENCRYPTED_DOCUMENT)
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _extract_docx(content: bytes) -> str:
    document = Document(BytesIO(content))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            paragraphs.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(paragraphs)


def _extract_with_libreoffice(filename: str, content: bytes) -> str:
    executable = shutil.which("soffice")
    if executable is None:
        raise DocumentExtractionError(CvImportErrorCode.EXTRACTOR_FAILED)
    safe_suffix = Path(filename).suffix.lower()
    with tempfile.TemporaryDirectory(prefix="careerops-cv-") as raw_dir:
        workdir = Path(raw_dir)
        source = workdir / f"source{safe_suffix}"
        source.write_bytes(content)
        profile = (workdir / "profile").as_uri()
        environment = os.environ.copy()
        completed = subprocess.run(
            [
                executable,
                "--headless",
                "--nologo",
                "--nodefault",
                "--nolockcheck",
                "--norestore",
                f"-env:UserInstallation={profile}",
                "--convert-to",
                "txt:Text",
                "--outdir",
                str(workdir),
                str(source),
            ],
            check=False,
            capture_output=True,
            timeout=30,
            env=environment,
        )
        output = workdir / "source.txt"
        if completed.returncode != 0 or not output.is_file():
            raise DocumentExtractionError(CvImportErrorCode.EXTRACTOR_FAILED)
        return output.read_text(encoding="utf-8-sig")


def _validate_archive(content: bytes) -> None:
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            members = archive.infolist()
            if len(members) > MAX_ARCHIVE_MEMBERS:
                raise DocumentExtractionError(CvImportErrorCode.INVALID_ARCHIVE)
            expanded = sum(member.file_size for member in members)
            compressed = sum(max(member.compress_size, 1) for member in members)
            if expanded > MAX_ARCHIVE_EXPANDED_BYTES or expanded / compressed > MAX_ARCHIVE_RATIO:
                raise DocumentExtractionError(CvImportErrorCode.INVALID_ARCHIVE)
    except zipfile.BadZipFile as exc:
        raise DocumentExtractionError(CvImportErrorCode.INVALID_ARCHIVE) from exc
